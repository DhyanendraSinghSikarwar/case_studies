"""Recalculate Excel formulas with headless LibreOffice and return the computed values.

Usage (from your own scratch folder):
    import sys; sys.path.insert(0, '/home/user/case_studies/Interview_DA/tools')
    from xl_eval import evaluate
    res = evaluate({'Sheet1': {'A1': 'Region', 'B1': 'Sales', 'A2': 'N', 'B2': 10, 'D1': '=SUM(B2:B3)'}})
    res['Sheet1']['D1']   # -> computed value

Rules
- Cells: plain values, or strings starting with '=' for formulas. Dates: pass datetime.date objects.
- Functions newer than Excel 2010 need the _xlfn. prefix when written through openpyxl, e.g.
  '=_xlfn.TEXTJOIN(",",TRUE,B1:B3)', '=_xlfn.IFS(...)', '=_xlfn.MAXIFS(...)', '=_xlfn.CONCAT(...)', '=_xlfn.SWITCH(...)',
  '=_xlfn.RANK.EQ(...)'. Verified working in LibreOffice 24.2: TEXTJOIN, IFS, MAXIFS, CONCAT, SWITCH, RANK.EQ,
  SUMIFS, COUNTIFS, AVERAGEIFS, INDEX/MATCH, VLOOKUP, HLOOKUP, SUMPRODUCT, DATEDIF, NETWORKDAYS, EOMONTH, TEXT.
- NOT supported by LibreOffice 24.2 (returns #NAME?/#VALUE!): XLOOKUP, XMATCH, FILTER, UNIQUE, SORT, SORTBY, LET,
  TEXTSPLIT, TEXTBEFORE/TEXTAFTER, dynamic arrays/spill, LAMBDA, Power Query, pivot tables, VBA, Copilot.
  For these, label the answer 'not run here' and reproduce the expected result with pandas or by hand.
- LibreOffice results can differ from Excel in edge cases (text-to-number coercion, locale, 1900 date quirks); say so
  if it matters.
"""
import os, shutil, subprocess, tempfile
import openpyxl


def evaluate(sheets, timeout=180):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, cells in sheets.items():
        ws = wb.create_sheet(name)
        for addr, v in cells.items():
            ws[addr] = v
    tmp = tempfile.mkdtemp(prefix='xl_')
    try:
        src = os.path.join(tmp, 'in.xlsx')
        wb.save(src)
        out = os.path.join(tmp, 'out')
        profile = 'file://' + os.path.join(tmp, 'profile')  # private profile => safe in parallel
        subprocess.run(['soffice', f'-env:UserInstallation={profile}', '--headless', '--calc',
                        '--convert-to', 'xlsx:Calc MS Excel 2007 XML', src, '--outdir', out],
                       check=True, capture_output=True, timeout=timeout)
        res = openpyxl.load_workbook(os.path.join(out, 'in.xlsx'), data_only=True)
        return {ws.title: {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None}
                for ws in res.worksheets}
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == '__main__':
    r = evaluate({'S': {'A1': 3, 'A2': 4, 'B1': '=SUM(A1:A2)', 'B2': '=_xlfn.TEXTJOIN(",",TRUE,A1:A2)'}})
    print(r)
