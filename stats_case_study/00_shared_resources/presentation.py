"""Shared python-pptx deck builder for the stakeholder presentations.

Design rules baked in (so every case's deck looks the same):
- 16:9 slides, one message per slide, the slide title states the "so what"
  (an action title, not a topic label);
- an optional takeaway bar at the bottom of a slide;
- no formulas; numbers are rounded and written in business units;
- footer with the case name and slide number.

Inline emphasis: wrap text in **double asterisks** to make it bold.
"""

from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

NAVY = RGBColor(0x14, 0x2B, 0x4C)
BLUE = RGBColor(0x2A, 0x78, 0xD6)
ORANGE = RGBColor(0xEB, 0x68, 0x34)
INK = RGBColor(0x1A, 0x1A, 0x1A)
INK_2 = RGBColor(0x52, 0x51, 0x4E)
MUTED = RGBColor(0x8A, 0x89, 0x85)
LIGHT_BG = RGBColor(0xEE, 0xF4, 0xFC)
TILE_BG = RGBColor(0xF5, 0xF5, 0xF3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.6)
FONT = "Calibri"


def _add_runs(paragraph, text: str, size: int, color: RGBColor = INK, bold_all: bool = False) -> None:
    """Add text to a paragraph, turning **segments** into bold runs.

    A hyphen used as a minus sign (e.g. "-1.2%") is replaced by a true minus ("−1.2%");
    hyphens inside words or ranges ("Tier-3", "40-60") are left alone.
    """
    text = re.sub(r"(?<![\w\-.])-(?=[\d₹])", "−", text)
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if not part:
            continue
        run = paragraph.add_run()
        is_bold = part.startswith("**") and part.endswith("**")
        run.text = part[2:-2] if is_bold else part
        run.font.size = Pt(size)
        run.font.name = FONT
        run.font.bold = bold_all or is_bold
        run.font.color.rgb = color


class Deck:
    """Thin wrapper around python-pptx with a fixed visual language."""

    def __init__(self, footer: str):
        self.prs = Presentation()
        self.prs.slide_width = SLIDE_W
        self.prs.slide_height = SLIDE_H
        self.footer = footer
        self._blank = self.prs.slide_layouts[6]

    # ------------------------------------------------------------------ basics
    def _new_slide(self, title: str | None):
        slide = self.prs.slides.add_slide(self._blank)
        if title:
            box = slide.shapes.add_textbox(MARGIN, Inches(0.35), SLIDE_W - 2 * MARGIN, Inches(1.0))
            tf = box.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.BOTTOM
            _add_runs(tf.paragraphs[0], title, 26, NAVY, bold_all=True)
            line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN, Inches(1.42), Inches(1.2), Inches(0.06))
            line.fill.solid()
            line.fill.fore_color.rgb = ORANGE
            line.line.fill.background()
        self._footer(slide)
        return slide

    def _footer(self, slide) -> None:
        n = len(self.prs.slides)
        box = slide.shapes.add_textbox(MARGIN, SLIDE_H - Inches(0.45), Inches(9), Inches(0.3))
        _add_runs(box.text_frame.paragraphs[0], self.footer, 10, MUTED)
        num = slide.shapes.add_textbox(SLIDE_W - MARGIN - Inches(1), SLIDE_H - Inches(0.45), Inches(1), Inches(0.3))
        p = num.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        _add_runs(p, str(n), 10, MUTED)

    def _takeaway(self, slide, text: str) -> None:
        top = SLIDE_H - Inches(1.45)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN, top, SLIDE_W - 2 * MARGIN, Inches(0.85))
        bar.fill.solid()
        bar.fill.fore_color.rgb = LIGHT_BG
        bar.line.fill.background()
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN, top, Inches(0.09), Inches(0.85))
        accent.fill.solid()
        accent.fill.fore_color.rgb = BLUE
        accent.line.fill.background()
        tf = bar.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Inches(0.35)
        tf.margin_right = Inches(0.3)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        _add_runs(p, text, 16, NAVY)

    def _bullets_box(self, slide, left, top, width, height, bullets, size: int = 18) -> None:
        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = True
        first = True
        for item in bullets:
            text, level = (item if isinstance(item, tuple) else (item, 0))
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.space_after = Pt(10 if level == 0 else 4)
            marker = "•  " if level == 0 else "–  "
            indent = "      " * level
            _add_runs(p, indent + marker, size - 2 * level, BLUE if level == 0 else MUTED, bold_all=True)
            _add_runs(p, text, size - 2 * level, INK if level == 0 else INK_2)

    def _body_bottom(self, takeaway: str | None):
        return SLIDE_H - (Inches(1.65) if takeaway else Inches(0.7))

    # ------------------------------------------------------------- slide types
    def title_slide(self, title: str, subtitle: str, note: str = "") -> None:
        slide = self.prs.slides.add_slide(self._blank)
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
        bg.fill.solid()
        bg.fill.fore_color.rgb = NAVY
        bg.line.fill.background()
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, MARGIN, Inches(3.55), Inches(1.6), Inches(0.08))
        bar.fill.solid()
        bar.fill.fore_color.rgb = ORANGE
        bar.line.fill.background()
        box = slide.shapes.add_textbox(MARGIN, Inches(1.5), SLIDE_W - 2 * MARGIN, Inches(2.0))
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.BOTTOM
        _add_runs(tf.paragraphs[0], title, 38, WHITE, bold_all=True)
        sub = slide.shapes.add_textbox(MARGIN, Inches(3.85), SLIDE_W - 2 * MARGIN, Inches(1.5))
        sub.text_frame.word_wrap = True
        _add_runs(sub.text_frame.paragraphs[0], subtitle, 20, RGBColor(0xC9, 0xD6, 0xEA))
        if note:
            nb = slide.shapes.add_textbox(MARGIN, SLIDE_H - Inches(1.0), SLIDE_W - 2 * MARGIN, Inches(0.5))
            nb.text_frame.word_wrap = True
            _add_runs(nb.text_frame.paragraphs[0], note, 12, RGBColor(0xA9, 0xB8, 0xCF))

    def bullets(self, title: str, bullets: list, takeaway: str | None = None, size: int = 20) -> None:
        slide = self._new_slide(title)
        top = Inches(1.75)
        self._bullets_box(slide, MARGIN, top, SLIDE_W - 2 * MARGIN, self._body_bottom(takeaway) - top, bullets, size)
        if takeaway:
            self._takeaway(slide, takeaway)

    def image(self, title: str, image_path: Path, bullets: list | None = None,
              takeaway: str | None = None, image_share: float = 0.62) -> None:
        """Chart slide. With bullets, the chart takes ``image_share`` of the width and bullets sit to the right."""
        slide = self._new_slide(title)
        top = Inches(1.7)
        bottom = self._body_bottom(takeaway)
        avail_w = SLIDE_W - 2 * MARGIN
        img_w_max = int(avail_w * image_share) if bullets else avail_w
        img_h_max = bottom - top
        with Image.open(image_path) as im:
            w_px, h_px = im.size
        scale = min(img_w_max / w_px, img_h_max / h_px)
        w, h = int(w_px * scale), int(h_px * scale)
        left = MARGIN if bullets else int((SLIDE_W - w) / 2)
        slide.shapes.add_picture(str(image_path), Emu(left), Emu(top + int((img_h_max - h) / 2)), Emu(w), Emu(h))
        if bullets:
            b_left = MARGIN + img_w_max + Inches(0.35)
            self._bullets_box(slide, b_left, top + Inches(0.1), SLIDE_W - MARGIN - b_left, img_h_max, bullets, 16)
        if takeaway:
            self._takeaway(slide, takeaway)

    def kpis(self, title: str, items: list[tuple[str, str, str]], takeaway: str | None = None) -> None:
        """Up to four headline numbers: (value, label, small explanatory note)."""
        slide = self._new_slide(title)
        n = len(items)
        gap = Inches(0.3)
        tile_w = int((SLIDE_W - 2 * MARGIN - gap * (n - 1)) / n)
        top, tile_h = Inches(2.0), Inches(3.3)
        for i, (value, label, note) in enumerate(items):
            left = MARGIN + i * (tile_w + gap)
            tile = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, tile_w, tile_h)
            tile.fill.solid()
            tile.fill.fore_color.rgb = TILE_BG
            tile.line.fill.background()
            tf = tile.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.TOP
            tf.margin_left = tf.margin_right = Inches(0.25)
            tf.margin_top = Inches(0.3)
            p = tf.paragraphs[0]
            _add_runs(p, value, 36, NAVY, bold_all=True)
            p2 = tf.add_paragraph()
            p2.space_before = Pt(6)
            _add_runs(p2, label, 16, INK, bold_all=True)
            p3 = tf.add_paragraph()
            p3.space_before = Pt(8)
            _add_runs(p3, note, 13, INK_2)
        if takeaway:
            self._takeaway(slide, takeaway)

    def table(self, title: str, df: pd.DataFrame, takeaway: str | None = None,
              col_widths: list[float] | None = None, font_size: int = 14, highlight_rows: list[int] | None = None) -> None:
        """Simple, readable table. ``df`` should already contain formatted strings."""
        slide = self._new_slide(title)
        rows, cols = df.shape[0] + 1, df.shape[1]
        top = Inches(1.8)
        width = SLIDE_W - 2 * MARGIN
        row_h = Inches(0.25 + font_size / 72 * 0.9)
        shape = slide.shapes.add_table(rows, cols, MARGIN, top, width, row_h * rows)
        tbl = shape.table
        if col_widths:
            total = sum(col_widths)
            for j, cw in enumerate(col_widths):
                tbl.columns[j].width = int(width * cw / total)
        highlight_rows = highlight_rows or []
        for j, col in enumerate(df.columns):
            cell = tbl.cell(0, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY
            cell.text_frame.paragraphs[0].text = ""
            _add_runs(cell.text_frame.paragraphs[0], str(col), font_size, WHITE, bold_all=True)
        for i, (_, row) in enumerate(df.iterrows(), start=1):
            for j, val in enumerate(row):
                cell = tbl.cell(i, j)
                cell.fill.solid()
                cell.fill.fore_color.rgb = LIGHT_BG if (i - 1) in highlight_rows else (WHITE if i % 2 else TILE_BG)
                cell.text_frame.paragraphs[0].text = ""
                cell.text_frame.word_wrap = True
                _add_runs(cell.text_frame.paragraphs[0], str(val), font_size, INK, bold_all=(j == 0))
        if takeaway:
            self._takeaway(slide, takeaway)

    def two_columns(self, title: str, left_title: str, left: list, right_title: str, right: list,
                    takeaway: str | None = None, size: int = 17) -> None:
        slide = self._new_slide(title)
        top = Inches(1.75)
        col_w = int((SLIDE_W - 2 * MARGIN - Inches(0.5)) / 2)
        bottom = self._body_bottom(takeaway)
        for k, (head, items) in enumerate([(left_title, left), (right_title, right)]):
            left_x = MARGIN + k * (col_w + Inches(0.5))
            hb = slide.shapes.add_textbox(left_x, top, col_w, Inches(0.5))
            _add_runs(hb.text_frame.paragraphs[0], head, 19, BLUE if k == 0 else ORANGE, bold_all=True)
            self._bullets_box(slide, left_x, top + Inches(0.6), col_w, bottom - top - Inches(0.6), items, size)
        if takeaway:
            self._takeaway(slide, takeaway)

    def save(self, path: Path) -> Path:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(path)
        return path
